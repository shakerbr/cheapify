"""Drawing helpers + animation XML engine (python-pptx / lxml)."""
import copy
import struct
import tempfile
import zipfile
from pathlib import Path
from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from PIL import Image

from assets import footer_icon_paths, icon_path

EMU = 914400

INK = "101820"
INK60 = "5A6672"
INK35 = "A9B1B9"
BLUE = "298FC2"
BLUE12 = "E5F2F8"
BLUE50 = "94C7E0"
GREEN = "228848"
GREEN8 = "EDF5F0"
RED = "D22730"
ORANGE = "FFA300"
NEUTRAL = "DDE5ED"
SURFACE = "F4F6F8"
WHITE = "FFFFFF"

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}
NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"


def E(v):
    return Emu(int(round(v * EMU)))


def rgb(hexstr):
    return RGBColor.from_string(hexstr)


# ---------------------------------------------------------------- shadows
SHADOWS = {
    "elev1": (8, 2, 9000),
    "elev2": (24, 8, 8000),
}


def no_effect(shape):
    spPr = shape._element.spPr
    for old in spPr.findall(qn("a:effectLst")):
        spPr.remove(old)
    etree.SubElement(spPr, qn("a:effectLst"))
    _drop_style(shape)


def _drop_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def add_shadow(shape, kind="elev1"):
    blur, dist, alpha = SHADOWS[kind]
    spPr = shape._element.spPr
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    _drop_style(shape)
    sh = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(blur * 12700),
                          dist=str(dist * 12700), dir="5400000", algn="t", rotWithShape="0")
    clr = etree.SubElement(sh, qn("a:srgbClr"), val=INK)
    etree.SubElement(clr, qn("a:alpha"), val=str(alpha))


# ---------------------------------------------------------------- shapes
def rect(parent, x, y, w, h, fill=None, line=None, lw=1.0, radius=None, shadow=None,
         shape=None, name=None, dash=None):
    if shape is None:
        shape = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sp = parent.add_shape(shape, E(x), E(y), E(w), E(h))
    if radius and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = min(0.5, radius / min(w, h))
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = rgb(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = rgb(line)
        sp.line.width = Pt(lw)
        if dash:
            sp.line.dash_style = dash
    else:
        sp.line.fill.background()
    if shadow:
        add_shadow(sp, shadow)
    else:
        no_effect(sp)
    if name:
        sp.name = name
    return sp


def oval(parent, cx, cy, r, fill=None, line=None, lw=1.0, name=None):
    return rect(parent, cx - r, cy - r, 2 * r, 2 * r, fill=fill, line=line, lw=lw,
                shape=MSO_SHAPE.OVAL, name=name)


def card(parent, x, y, w, h, shadow="elev1", fill=WHITE, line=NEUTRAL, lw=1.0, radius=0.14):
    return rect(parent, x, y, w, h, fill=fill, line=line, lw=lw, radius=radius, shadow=shadow)


def line(parent, x1, y1, x2, y2, color=INK35, lw=1.0, dash=None):
    c = parent.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(lw)
    if dash:
        c.line.dash_style = dash
    no_effect(c)
    return c


def group(parent):
    new_id = max(parent._spTree.max_shape_id + 1, 2)
    g = parent.add_group_shape()
    cNvPr = g._element.find(qn("p:nvGrpSpPr") + "/" + qn("p:cNvPr"))
    cNvPr.set("id", str(new_id))
    cNvPr.set("name", f"Group {new_id}")
    return g


def _set_font(rPr, face, rtl=False):
    for tag in ("a:latin", "a:ea", "a:cs", "a:sym"):
        for old in rPr.findall(qn(tag)):
            rPr.remove(old)
    for tag in ("a:latin", "a:cs"):
        etree.SubElement(rPr, qn(tag), typeface=face)


def text(parent, x, y, w, h, paras, size=12, color=INK, bold=False, font="Inter",
         align="l", anchor="t", rtl=False, wrap=True, spacing=None, name=None):
    """paras: str | [para]; para: str | [run]; run: str | (str, {overrides})."""
    tb = parent.add_textbox(E(x), E(y), E(w), E(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR[anchor]
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for ch in list(bodyPr):
        bodyPr.remove(ch)
    etree.SubElement(bodyPr, qn("a:noAutofit"))
    if isinstance(paras, str):
        paras = [paras]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN[align]
        if spacing:
            p.line_spacing = spacing
        if rtl:
            p._p.get_or_add_pPr().set("rtl", "1")
        runs = [para] if isinstance(para, (str, tuple)) else para
        for run in runs:
            if isinstance(run, str):
                t, o = run, {}
            else:
                t, o = run
            r = p.add_run()
            r.text = t
            f = r.font
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = rgb(o.get("color", color))
            rPr = r._r.get_or_add_rPr()
            fnt = o.get("font", font)
            rPr.set("lang", "ar-IQ" if rtl else "en-US")
            _set_font(rPr, fnt, rtl)
    return tb


def pic_contain(parent, path, x, y, w, h, align="c"):
    im = Image.open(path)
    iw, ih = im.size
    s = min(w / iw, h / ih)
    pw, ph = iw * s, ih * s
    px = x + (w - pw) / 2
    py = y + (h - ph) / 2
    pic = parent.add_picture(str(path), E(px), E(py), E(pw), E(ph))
    return pic, (px, py, s)


def pic_box(parent, path, x, y, w, h):
    return parent.add_picture(str(path), E(x), E(y), E(w), E(h))


def icon(parent, name, color, x, y, s):
    return parent.add_picture(str(icon_path(name, color)), E(x), E(y), E(s), E(s))


def icon_c(parent, name, color, cx, cy, s):
    return icon(parent, name, color, cx - s / 2, cy - s / 2, s)


def ring(parent, pic_geom, src_box, pad=6, color=RED, lw=2.25):
    px, py, s = pic_geom
    x0, y0, x1, y1 = src_box
    return rect(parent, px + (x0 - pad) * s, py + (y0 - pad) * s,
                (x1 - x0 + 2 * pad) * s, (y1 - y0 + 2 * pad) * s,
                line=color, lw=lw, radius=0.06)


# ---------------------------------------------------------------- template bits
NS_ASVG = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"


def resize_headers(slide, cover=False):
    """Set header logo sizes/positions to match archive/GIZ-assets/template.jpg (v4 spec §1) while preserving exact native image aspect ratios, normalizing spTree ID=1, and upgrading footer icons to SVG on cover/thanks."""
    # Ensure root spTree has id="1" (PowerPoint reserves id=1 for spTree/bg)
    spTree = slide.shapes._spTree
    root_cNvPr = spTree.find(qn("p:nvGrpSpPr") + "/" + qn("p:cNvPr"))
    if root_cNvPr is not None:
        root_cNvPr.set("id", "1")
        root_cNvPr.set("name", "")
    if not cover:
        cSld = slide._element.find(qn("p:cSld"))
        if cSld is not None:
            for bg in cSld.findall(qn("p:bg")):
                cSld.remove(bg)

    # Cover/thanks: German Cooperation w=2.60 at (0.47, 0.40); CAPA w=1.76 at (11.12, 1.20)
    # Content (75%): German Cooperation w=1.95 at (0.47, 0.30); CAPA w=1.32 at (11.56, 0.90)
    giz = (0.47, 0.40, 2.60) if cover else (0.47, 0.30, 1.95)
    capa = (12.88 - 1.76, 1.20, 1.76) if cover else (12.88 - 1.32, 0.90, 1.32)
    old_footer_pics = []
    for sh in slide.shapes:
        if sh.shape_type == 13 and sh.top < E(2.0):
            x, y, w = giz if sh.left < E(5.0) else capa
            iw, ih = sh.image.size
            h = w * ih / iw
            sh.left, sh.top, sh.width, sh.height = E(x), E(y), E(w), E(h)
            cNvPicPr = sh._element.find(qn("p:nvPicPr") + "/" + qn("p:cNvPicPr"))
            if cNvPicPr is not None:
                cNvPicPr.set("preferRelativeResize", "1")
        elif cover and sh.shape_type == 13 and sh.top > E(5.0):
            old_footer_pics.append(sh)
        elif cover and sh.has_text_frame and sh.top > E(5.0):
            txt = sh.text_frame.text.strip()
            if txt == "capa.auk":
                sh.left, sh.top = E(1.844), E(6.573)
            elif txt == "CAPA AUK":
                sh.left, sh.top = E(6.026), E(6.573)
            elif txt == "capa@auk.edu.krd":
                sh.left, sh.top = E(10.323), E(6.573)

    if cover and old_footer_pics:
        for sh in old_footer_pics:
            spTree.remove(sh._element)
        footer_specs = [
            (1.167, "footer-facebook"),
            (1.485, "footer-instagram"),
            (5.680, "footer-linkedin"),
            (9.969, "footer-mail"),
        ]
        for fx_in, icon_name in footer_specs:
            _, png_path = footer_icon_paths(icon_name)
            pic = slide.shapes.add_picture(str(png_path), E(fx_in), E(6.557), E(0.24), E(0.24))
            pic._element.set("data-svg-icon", icon_name)


def clone_header(src_slide, dst_slide):
    """Copy the two header-logo pictures (same image parts) from src to dst."""
    for pic in src_slide.shapes._spTree.findall(qn("p:pic")):
        new = copy.deepcopy(pic)
        blip = new.find(".//" + qn("a:blip"))
        rid = blip.get(qn("r:embed"))
        img_part = src_slide.part.related_part(rid)
        new_rid = dst_slide.part.relate_to(img_part, RT.IMAGE)
        blip.set(qn("r:embed"), new_rid)
        dst_slide.shapes._spTree.append(new)


def content_frame(slide, label, headline):
    n = len(headline)
    hsz = 30 if n <= 44 else (25 if n <= 52 else (22 if n <= 62 else 19.5))
    b = rect(slide.shapes, 0.79, 1.66, 0.30, 0.05, fill=BLUE)
    l = text(slide.shapes, 1.24, 1.56, 6.0, 0.28, label, size=13, color=INK60, anchor="m")
    h = text(slide.shapes, 0.79, 1.86, 11.75, 0.62, headline, size=hsz, color=INK, bold=True,
             font="Comfortaa", anchor="m", wrap=False)
    return b, l, h


def footnote(slide, txt):
    return text(slide.shapes, 0.79, 7.06, 11.75, 0.26, txt, size=9, color=INK60, anchor="m")


# ---------------------------------------------------------------- animation engine
class Fx:
    def __init__(self, shape, kind="fade", delay=0, dur=None):
        self.shape = shape
        self.kind = kind
        self.delay = delay
        self.dur = dur if dur is not None else {"fade": 300, "zoom": 300, "appear": 1,
                                                  "fadeout": 250}[kind]


def fx(shape, kind="fade", delay=0, dur=None):
    return Fx(shape, kind, delay, dur)


def stagger(shapes, kind="fade", gap=None, dur=None, start=0):
    """One step of withEffects with cumulative delays (4+ items: 220 ms / 40 ms; else 280 / 90)."""
    many = len(shapes) >= 4
    gap = gap if gap is not None else (40 if many else 90)
    dur = dur if dur is not None else (220 if many else 280)
    if kind == "appear":
        dur = 1
    return [[fx(s, kind, start + i * gap, dur) for i, s in enumerate(shapes)]]


def pair_stagger(pairs, kind="fade", gap=None, dur=None):
    """Like stagger() but each item is a list of shapes sharing one delay."""
    many = len(pairs) >= 4
    gap = gap if gap is not None else (40 if many else 90)
    dur = dur if dur is not None else (220 if many else 280)
    step = []
    for i, grp in enumerate(pairs):
        for s in grp:
            step.append(fx(s, kind, i * gap, dur))
    return [step]


def together(*effects):
    return [list(effects)]


class Anim:
    """Collects click groups; group = list of steps; step = list[Fx]."""

    def __init__(self):
        self.groups = []

    def add(self, steps, auto=False):
        if steps and isinstance(steps[0], Fx):
            steps = [steps]
        self.groups.append((auto, [s for s in steps if s]))

    def apply_state(self, upto):
        """QA: remove shapes whose entrance happens after click `upto` and exits up to it."""
        for gi, (auto, steps) in enumerate(self.groups):
            for step in steps:
                for e in step:
                    el = e.shape._element
                    parent = el.getparent()
                    if parent is None:
                        continue
                    if e.kind == "fadeout":
                        if gi <= upto:
                            parent.remove(el)
                    elif gi > upto:
                        parent.remove(el)

    def build(self, slide):
        if not self.groups:
            return
        nid = [2]

        def nxt():
            nid[0] += 1
            return nid[0]

        used = {}
        bld = []
        P = "{%s}" % NS_P
        timing = etree.Element(P + "timing")
        tnLst = etree.SubElement(timing, P + "tnLst")
        par0 = etree.SubElement(tnLst, P + "par")
        root = etree.SubElement(par0, P + "cTn", id="1", dur="indefinite", restart="never",
                                nodeType="tmRoot")
        rc = etree.SubElement(root, P + "childTnLst")
        seq = etree.SubElement(rc, P + "seq", concurrent="1", nextAc="seek")
        main = etree.SubElement(seq, P + "cTn", id="2", dur="indefinite", nodeType="mainSeq")
        mc = etree.SubElement(main, P + "childTnLst")
        prev = etree.SubElement(seq, P + "prevCondLst")
        c = etree.SubElement(prev, P + "cond", evt="onPrev", delay="0")
        etree.SubElement(etree.SubElement(c, P + "tgtEl"), P + "sldTgt")
        nxtl = etree.SubElement(seq, P + "nextCondLst")
        c = etree.SubElement(nxtl, P + "cond", evt="onNext", delay="0")
        etree.SubElement(etree.SubElement(c, P + "tgtEl"), P + "sldTgt")

        for auto, steps in self.groups:
            gpar = etree.SubElement(mc, P + "par")
            gct = etree.SubElement(gpar, P + "cTn", id=str(nxt()), fill="hold")
            st = etree.SubElement(gct, P + "stCondLst")
            etree.SubElement(st, P + "cond", delay="indefinite")
            if auto:
                etree.SubElement(st, P + "cond", evt="onBegin", delay="0").append(
                    etree.Element(P + "tn", val="2"))
            gcl = etree.SubElement(gct, P + "childTnLst")
            start = 0
            for si, step in enumerate(steps):
                spar = etree.SubElement(gcl, P + "par")
                sct = etree.SubElement(spar, P + "cTn", id=str(nxt()), fill="hold")
                sst = etree.SubElement(sct, P + "stCondLst")
                etree.SubElement(sst, P + "cond", delay=str(start))
                scl = etree.SubElement(sct, P + "childTnLst")
                for ei, e in enumerate(step):
                    if ei == 0:
                        node = "clickEffect" if (si == 0 and not auto) else "afterEffect"
                    else:
                        node = "withEffect"
                    spid = e.shape.shape_id
                    grp = used.get(spid, 0)
                    used[spid] = grp + 1
                    if e.shape._element.tag == qn("p:sp"):
                        bld.append((spid, grp))
                    self._effect(scl, e, node, grp, nxt)
                start += max(x.delay + x.dur for x in step)

        if bld:
            bl = etree.SubElement(timing, P + "bldLst")
            for spid, grp in bld:
                etree.SubElement(bl, P + "bldP", spid=str(spid), grpId=str(grp), animBg="1")

        cSld_parent = slide._element
        for old in cSld_parent.findall(qn("p:timing")):
            cSld_parent.remove(old)
        cSld_parent.append(timing)

    @staticmethod
    def _effect(parent, e, node, grp, nxt):
        P = "{%s}" % NS_P
        exit_ = e.kind == "fadeout"
        preset = {"fade": ("10", "0"), "zoom": ("53", "16"), "appear": ("1", "0"),
                  "fadeout": ("10", "0")}[e.kind]
        par = etree.SubElement(parent, P + "par")
        ct = etree.SubElement(par, P + "cTn", id=str(nxt()), presetID=preset[0],
                              presetClass="exit" if exit_ else "entr",
                              presetSubtype=preset[1], fill="hold", grpId=str(grp),
                              nodeType=node)
        st = etree.SubElement(ct, P + "stCondLst")
        etree.SubElement(st, P + "cond", delay=str(e.delay))
        cl = etree.SubElement(ct, P + "childTnLst")
        spid = str(e.shape.shape_id)

        def tgt(bhvr):
            etree.SubElement(etree.SubElement(bhvr, P + "tgtEl"), P + "spTgt", spid=spid)

        def set_vis(val, delay="0"):
            s = etree.SubElement(cl, P + "set")
            b = etree.SubElement(s, P + "cBhvr")
            c2 = etree.SubElement(b, P + "cTn", id=str(nxt()), dur="1", fill="hold")
            etree.SubElement(etree.SubElement(c2, P + "stCondLst"), P + "cond", delay=delay)
            tgt(b)
            an = etree.SubElement(b, P + "attrNameLst")
            etree.SubElement(an, P + "attrName").text = "style.visibility"
            to = etree.SubElement(s, P + "to")
            etree.SubElement(to, P + "strVal", val=val)

        def fade(trans):
            a = etree.SubElement(cl, P + "animEffect", transition=trans, filter="fade")
            b = etree.SubElement(a, P + "cBhvr")
            etree.SubElement(b, P + "cTn", id=str(nxt()), dur=str(e.dur))
            tgt(b)

        if e.kind == "appear":
            set_vis("visible")
        elif e.kind == "fade":
            set_vis("visible")
            fade("in")
        elif e.kind == "zoom":
            set_vis("visible")
            for attr in ("ppt_w", "ppt_h"):
                a = etree.SubElement(cl, P + "anim", calcmode="lin", valueType="num")
                b = etree.SubElement(a, P + "cBhvr")
                etree.SubElement(b, P + "cTn", id=str(nxt()), dur=str(e.dur), fill="hold")
                tgt(b)
                an = etree.SubElement(b, P + "attrNameLst")
                etree.SubElement(an, P + "attrName").text = attr
                tl = etree.SubElement(a, P + "tavLst")
                t0 = etree.SubElement(tl, P + "tav", tm="0")
                etree.SubElement(etree.SubElement(t0, P + "val"), P + "fltVal", val="0")
                t1 = etree.SubElement(tl, P + "tav", tm="100000")
                etree.SubElement(etree.SubElement(t1, P + "val"), P + "strVal", val="#" + attr)
            fade("in")
        elif e.kind == "fadeout":
            fade("out")
            set_vis("hidden", delay=str(e.dur - 1))


def finalize_slide(slide, anim=None, transition=True):
    root = slide._element
    for old in root.findall(qn("p:transition")):
        root.remove(old)
    if transition:
        P = "{%s}" % NS_P
        tr = etree.Element(P + "transition", spd="fast")
        etree.SubElement(tr, P + "fade")
        clr = root.find(qn("p:clrMapOvr"))
        clr.addnext(tr)
    if anim is not None:
        anim.build(slide)


# ---------------------------------------------------------------- arcs / gauges
def _geom_adj(shape, vals):
    av = shape._element.spPr.find(qn("a:prstGeom")).find(qn("a:avLst"))
    for ch in list(av):
        av.remove(ch)
    for name, val in vals:
        etree.SubElement(av, qn("a:gd"), name=name, fmla=f"val {int(val)}")


def arc(parent, cx, cy, ro, ri, a0, a1, fill):
    sp = parent.add_shape(MSO_SHAPE.BLOCK_ARC, E(cx - ro), E(cy - ro), E(2 * ro), E(2 * ro))
    _geom_adj(sp, (("adj1", (a0 % 360) * 60000), ("adj2", (a1 % 360) * 60000),
                   ("adj3", (ro - ri) / (2 * ro) * 100000)))
    sp.fill.solid()
    sp.fill.fore_color.rgb = rgb(fill)
    sp.line.fill.background()
    no_effect(sp)
    return sp


def donut(parent, cx, cy, ro, ri, fill):
    sp = parent.add_shape(MSO_SHAPE.DONUT, E(cx - ro), E(cy - ro), E(2 * ro), E(2 * ro))
    _geom_adj(sp, (("adj", (ro - ri) / (2 * ro) * 100000),))
    sp.fill.solid()
    sp.fill.fore_color.rgb = rgb(fill)
    sp.line.fill.background()
    no_effect(sp)
    return sp


def pie_slice(parent, cx, cy, r, a0, a1, fill):
    sp = parent.add_shape(MSO_SHAPE.PIE, E(cx - r), E(cy - r), E(2 * r), E(2 * r))
    _geom_adj(sp, (("adj1", (a0 % 360) * 60000), ("adj2", (a1 % 360) * 60000)))
    sp.fill.solid()
    sp.fill.fore_color.rgb = rgb(fill)
    sp.line.fill.background()
    no_effect(sp)
    return sp


def gauge(parent, cx, cy, d, pct, color=BLUE, track=NEUTRAL, thick=0.14):
    ro = d / 2
    donut(parent, cx, cy, ro, ro - thick, track)
    arc(parent, cx, cy, ro, ro - thick, 270, 270 + 360 * pct, color)


# ---------------------------------------------------------------- font embedding
EMBED_FONTS = (
    ("Comfortaa", "Comfortaa-regular.fntdata", "Comfortaa-Regular.ttf",
     "Comfortaa-bold.fntdata", "Comfortaa-Bold.ttf"),
    ("Inter", "Inter-regular.fntdata", "Inter-Regular.ttf",
     "Inter-bold.fntdata", "Inter-Bold.ttf"),
    ("IBM Plex Sans Arabic", "IBMPlexSansArabic-regular.fntdata", "IBMPlexSansArabic-Regular.ttf",
     "IBMPlexSansArabic-bold.fntdata", "IBMPlexSansArabic-Bold.ttf"),
)
NS_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"
FONT_REL_TYPE = f"{NS_R}/font"


def ttf_to_eot(ttf: bytes) -> bytes:
    """Wrap raw TrueType bytes in an uncompressed EOT v2.02 (.fntdata) header."""
    num_tables = struct.unpack(">H", ttf[4:6])[0]
    tables = {}
    for i in range(num_tables):
        off = 12 + i * 16
        tag, _, toff, tlen = struct.unpack(">4sIII", ttf[off:off + 16])
        tables[tag.decode("latin1")] = ttf[toff:toff + tlen]
    os2, head, name = tables["OS/2"], tables["head"], tables["name"]
    weight = struct.unpack(">H", os2[4:6])[0]
    panose = os2[32:42]
    ur = struct.unpack(">4I", os2[42:58])
    italic = 1 if (struct.unpack(">H", os2[62:64])[0] & 1) else 0
    cpr = struct.unpack(">2I", os2[78:86]) if len(os2) >= 86 else (1, 0)
    csa = struct.unpack(">I", head[8:12])[0]
    ncount, s_off = struct.unpack(">HH", name[2:6])
    names = {}
    for i in range(ncount):
        pid, _, lid, nid, slen, soff = struct.unpack(">6H", name[6 + i * 12:18 + i * 12])
        raw = name[s_off + soff:s_off + soff + slen]
        if pid == 3 and (nid not in names or lid == 0x0409):
            names[nid] = raw.decode("utf-16-be", errors="replace")
        elif pid == 1 and nid not in names:
            names[nid] = raw.decode("mac_roman", errors="replace")
    fam = names.get(1, "")
    sub = names.get(2, "Regular")
    ver = names.get(5, "Version 1.0")
    full = names.get(4, f"{fam} {sub}".strip())
    var_parts = []
    for s in (fam, sub, ver, full):
        b = s.encode("utf-16-le")
        var_parts.append(struct.pack("<HH", 0, len(b)) + b)
    var_parts.append(struct.pack("<HH", 0, 0))
    var_parts.append(b"BSGP" + struct.pack("<IHHII", 0, 0, 0, 0, 0))
    var_blob = b"".join(var_parts)
    font_size = len(ttf)
    eot_size = 80 + len(var_blob) + font_size
    fixed = struct.pack(
        "<IIII10sBBIHH4I2I5I",
        eot_size, font_size, 0x00020002, 0,
        panose, 1, italic, weight, 0, 0x504C,
        *ur, *cpr, csa, 0, 0, 0, 0,
    )
    return fixed + var_blob + ttf


def embed_fonts(pptx_path, font_dir, template_pptx=None):
    """Embed Comfortaa, Inter, and IBM Plex Sans Arabic (Regular + Bold) into pptx_path."""
    pptx_path = Path(pptx_path)
    font_dir = Path(font_dir)
    tpl_fonts = {}
    if template_pptx and Path(template_pptx).exists():
        with zipfile.ZipFile(template_pptx) as tz:
            for n in tz.namelist():
                if n.startswith("ppt/fonts/") and n.endswith(".fntdata"):
                    tpl_fonts[Path(n).name] = tz.read(n)

    fnt_blobs = {}
    for _, reg_fnt, reg_ttf, bold_fnt, bold_ttf in EMBED_FONTS:
        for fnt_name, ttf_name in ((reg_fnt, reg_ttf), (bold_fnt, bold_ttf)):
            if fnt_name in tpl_fonts:
                fnt_blobs[f"ppt/fonts/{fnt_name}"] = tpl_fonts[fnt_name]
            else:
                fnt_blobs[f"ppt/fonts/{fnt_name}"] = ttf_to_eot((font_dir / ttf_name).read_bytes())

    with zipfile.ZipFile(pptx_path, "r") as zin:
        entries = {name: zin.read(name) for name in zin.namelist()}

    # 1. [Content_Types].xml
    ct_root = etree.fromstring(entries["[Content_Types].xml"])
    if not any(el.get("Extension") == "fntdata" for el in ct_root.findall(f"{{{NS_CT}}}Default")):
        d_el = etree.Element(f"{{{NS_CT}}}Default", Extension="fntdata",
                             ContentType="application/x-fontdata")
        ct_root.insert(0, d_el)
    if not any(el.get("Extension") == "svg" for el in ct_root.findall(f"{{{NS_CT}}}Default")):
        s_el = etree.Element(f"{{{NS_CT}}}Default", Extension="svg",
                             ContentType="image/svg+xml")
        ct_root.insert(0, s_el)
    entries["[Content_Types].xml"] = etree.tostring(ct_root, xml_declaration=True,
                                                    encoding="UTF-8", standalone="yes")

    # 1b. Embed vector SVG blips for any <p:pic data-svg-icon="...">
    for sname in [n for n in entries if n.startswith("ppt/slides/slide") and n.endswith(".xml")]:
        s_root = etree.fromstring(entries[sname])
        svg_pics = [p for p in s_root.findall(".//" + qn("p:pic")) if p.get("data-svg-icon")]
        if not svg_pics:
            continue
        rels_path = f"ppt/slides/_rels/{Path(sname).name}.rels"
        sr_root = etree.fromstring(entries[rels_path])
        s_used = {r.get("Id") for r in sr_root.findall(f"{{{NS_REL}}}Relationship")}
        s_by_tgt = {r.get("Target"): r.get("Id") for r in sr_root.findall(f"{{{NS_REL}}}Relationship")}
        for pic_el in svg_pics:
            ic_name = pic_el.attrib.pop("data-svg-icon")
            svg_path, _ = footer_icon_paths(ic_name)
            media_name = f"ppt/media/{ic_name}.svg"
            entries[media_name] = svg_path.read_bytes()
            rel_tgt = f"../media/{ic_name}.svg"
            if rel_tgt in s_by_tgt:
                svg_rid = s_by_tgt[rel_tgt]
            else:
                idx = 1
                while f"rId{idx}" in s_used:
                    idx += 1
                svg_rid = f"rId{idx}"
                s_used.add(svg_rid)
                s_by_tgt[rel_tgt] = svg_rid
                etree.SubElement(sr_root, f"{{{NS_REL}}}Relationship",
                                 Id=svg_rid, Type=RT.IMAGE, Target=rel_tgt)
            blip = pic_el.find(".//" + qn("a:blip"))
            if blip is not None:
                ext_lst = blip.find(qn("a:extLst"))
                if ext_lst is None:
                    ext_lst = etree.SubElement(blip, qn("a:extLst"))
                ext = etree.SubElement(ext_lst, qn("a:ext"), uri=SVG_EXT_URI)
                svg_blip = etree.SubElement(
                    ext, f"{{{NS_ASVG}}}svgBlip",
                    nsmap={"asvg": NS_ASVG},
                )
                svg_blip.set(f"{{{NS_R}}}embed", svg_rid)
        entries[rels_path] = etree.tostring(sr_root, xml_declaration=True,
                                            encoding="UTF-8", standalone="yes")
        entries[sname] = etree.tostring(s_root, xml_declaration=True,
                                        encoding="UTF-8", standalone="yes")

    # 2. ppt/_rels/presentation.xml.rels
    rels_root = etree.fromstring(entries["ppt/_rels/presentation.xml.rels"])
    existing_by_target = {}
    used_ids = set()
    for rel in rels_root.findall(f"{{{NS_REL}}}Relationship"):
        used_ids.add(rel.get("Id"))
        existing_by_target[rel.get("Target")] = rel.get("Id")

    def ensure_font_rel(fnt_name):
        target = f"fonts/{fnt_name}"
        if target in existing_by_target:
            return existing_by_target[target]
        idx = 1
        while f"rId{idx}" in used_ids:
            idx += 1
        rid = f"rId{idx}"
        used_ids.add(rid)
        existing_by_target[target] = rid
        etree.SubElement(rels_root, f"{{{NS_REL}}}Relationship",
                         Id=rid, Type=FONT_REL_TYPE, Target=target)
        return rid

    rids = {}
    for fam, reg_fnt, _, bold_fnt, _ in EMBED_FONTS:
        rids[fam] = (ensure_font_rel(reg_fnt), ensure_font_rel(bold_fnt))
    entries["ppt/_rels/presentation.xml.rels"] = etree.tostring(
        rels_root, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )

    # 3. ppt/presentation.xml
    pres_root = etree.fromstring(entries["ppt/presentation.xml"])
    pres_root.set("embedTrueTypeFonts", "1")
    for old in pres_root.findall(f"{{{NS_P}}}embeddedFontLst"):
        pres_root.remove(old)
    ef_lst = etree.Element(f"{{{NS_P}}}embeddedFontLst")
    for fam, _, _, _, _ in EMBED_FONTS:
        reg_rid, bold_rid = rids[fam]
        ef = etree.SubElement(ef_lst, f"{{{NS_P}}}embeddedFont")
        etree.SubElement(ef, f"{{{NS_P}}}font", typeface=fam)
        etree.SubElement(ef, f"{{{NS_P}}}regular", {f"{{{NS_R}}}id": reg_rid})
        etree.SubElement(ef, f"{{{NS_P}}}bold", {f"{{{NS_R}}}id": bold_rid})

    notes_sz = pres_root.find(f"{{{NS_P}}}notesSz")
    if notes_sz is not None:
        notes_sz.addnext(ef_lst)
    else:
        pres_root.append(ef_lst)
    entries["ppt/presentation.xml"] = etree.tostring(
        pres_root, xml_declaration=True, encoding="UTF-8", standalone="yes"
    )

    # 4. Font binary parts
    entries.update(fnt_blobs)

    with tempfile.NamedTemporaryFile(dir=str(pptx_path.parent), suffix=".pptx", delete=False) as tf:
        tmp_path = Path(tf.name)
    with zipfile.ZipFile(tmp_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries.items():
            zout.writestr(name, data)
    tmp_path.replace(pptx_path)

