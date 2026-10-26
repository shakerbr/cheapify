#!/usr/bin/env python3
"""Rebuild the Cheapify GIZ pitch deck.

    python build/build_deck.py            # animated PPTX + static PDF + speaker script
    python build/build_deck.py --static   # only the static (no animation) PPTX
    python build/build_deck.py --qa-states  # also render intermediate-state slides

Run from the project root with the project venv (see presentation/build/README.md):
    tmp/venv/bin/python presentation/build/build_deck.py
QA renders go to tmp/deck-qa/ inside the project.
"""
import argparse
import copy
import re
import shutil
import subprocess
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from pptx import Presentation  # noqa: E402
from pptx.opc.constants import RELATIONSHIP_TYPE as RT  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

import assets  # noqa: E402
import deck_data as D  # noqa: E402
import slides_a  # noqa: E402
import slides_b  # noqa: E402
from deckkit import content_frame, finalize_slide, footnote  # noqa: E402

ROOT = HERE.parents[1]
PRES = ROOT / "presentation"
TEMPLATE = ROOT / "archive/program-giz-cheapify/Presentation-stuff/Modified template.pptx"
OUT_PPTX = PRES / "Cheapify-GIZ-Pitch.pptx"
OUT_PDF = PRES / "Cheapify-GIZ-Pitch.pdf"
OUT_GUIDE = PRES / "presenter-guide.md"
WPM = 140

MAIN = ["cover", "problem", "evidence", "solution", "value", "market", "customers", "product",
        "model", "competitors", "growth", "finance", "funding", "next", "thanks"]
APPX = ["A1", "A2", "A3", "A4", "A5", "A6", "A7", "A8"]


class Ctx:
    def __init__(self, static):
        self.static = static


def snapshot_header(slide):
    out = []
    for pic in slide.shapes._spTree.findall(qn("p:pic")):
        rid = pic.find(".//" + qn("a:blip")).get(qn("r:embed"))
        out.append((copy.deepcopy(pic), slide.part.related_part(rid)))
    return out


def add_header(slide, header):
    for pic, part in header:
        new = copy.deepcopy(pic)
        rid = slide.part.relate_to(part, RT.IMAGE)
        new.find(".//" + qn("a:blip")).set(qn("r:embed"), rid)
        slide.shapes._spTree.append(new)


def notes_text(key):
    d = D.SLIDES[key]
    txt = f"Time: ~{d['seconds']}s\n{d['notes']}"
    if d["keynums"]:
        txt += "\n\nKey numbers:\n" + "\n".join(f"{n} — {D.SRC[src]}" for n, _, src in d["keynums"])
    return txt


def set_notes(slide, key):
    slide.notes_slide.notes_text_frame.text = notes_text(key)


def build(static=False, qa_states=False, out=OUT_PPTX):
    meta = assets.build_all()
    slides_a.ASSETS = assets.OUT
    prs = Presentation(str(TEMPLATE))
    s_cover, s_content, s_thanks = list(prs.slides)
    layout = s_content.slide_layout
    header = snapshot_header(s_content)
    ctx = Ctx(static)
    order = {}
    anims = {}

    def content(key, reuse=None, label=None, headline=None, foot=None):
        if key in D.SLIDES:
            d = D.SLIDES[key]
            label, headline, foot = d["label"], d["headline"], d.get("footnote")
        sl = reuse if reuse is not None else prs.slides.add_slide(layout)
        if reuse is None:
            add_header(sl, header)
        content_frame(sl, label, headline)
        if foot:
            footnote(sl, foot)
        return sl

    # main deck
    order["cover"] = s_cover
    anims["cover"] = slides_a.cover(s_cover, ctx)
    order["problem"] = content("problem", reuse=s_content)
    anims["problem"] = slides_a.problem(order["problem"], ctx)
    for key, fn in (("evidence", slides_a.evidence), ("solution", None), ("value", slides_a.value),
                    ("market", slides_a.market), ("customers", slides_a.customers),
                    ("product", None), ("model", None),
                    ("competitors", slides_b.competitors), ("growth", slides_b.growth),
                    ("finance", slides_b.finance), ("funding", slides_b.funding),
                    ("next", slides_b.next_steps)):
        sl = content(key)
        order[key] = sl
        if key == "solution":
            anims[key] = slides_a.solution(sl, ctx, "static" if static else "anim")
        elif key == "product":
            anims[key] = slides_a.product(sl, ctx, "static" if static else "anim")
        elif key == "model":
            anims[key] = slides_b.model(sl, ctx, "static" if static else "anim")
        else:
            anims[key] = fn(sl, ctx)
    order["thanks"] = s_thanks
    anims["thanks"] = slides_b.thanks(s_thanks, ctx)

    # appendix
    for k in APPX:
        sl = content(k, label="Appendix", headline=D.APPENDIX[k]["headline"])
        order[k] = sl
        anims[k] = slides_b.appendix(sl, k, ctx)

    # QA-only intermediate states: (slide key, label, last click shown)
    extra = []
    if qa_states:
        specs = [("market", f"after click {k}", k) for k in (0, 1, 2)] + [
            ("model", "after click 4", 4), ("model", "after click 5", 5)] + [
            ("growth", f"after click {k}", k) for k in (0, 1, 2)] + [
            ("finance", "after click 1", 1), ("finance", "after click 2", 2)]
        builders = {"market": slides_a.market, "model": slides_b.model,
                    "growth": slides_b.growth, "finance": slides_b.finance}
        for key, lab, upto in specs:
            d = D.SLIDES[key]
            sl = content(f"qa_{key}_{lab}", label=d["label"], headline=d["headline"],
                         foot=d.get("footnote"))
            an = builders[key](sl, ctx, "anim") if key == "model" else builders[key](sl, ctx)
            an.apply_state(upto)
            extra.append(sl)

    # notes, transitions, animations
    for key in MAIN:
        set_notes(order[key], key)
    for k in APPX:
        order[k].notes_slide.notes_text_frame.text = D.APPENDIX_NOTE
    for key, sl in order.items():
        finalize_slide(sl, None if static else anims.get(key))
    for sl in extra:
        finalize_slide(sl, None)

    # slide order
    seq = [order[k] for k in MAIN + APPX] + extra
    lst = prs.slides._sldIdLst
    by_part = {}
    for sldId in list(lst):
        by_part[prs.part.related_part(sldId.rId)] = sldId
    for sldId in list(lst):
        lst.remove(sldId)
    for sl in seq:
        lst.append(by_part[sl.part])
    prs.save(str(out))
    return out


def convert_pdf(pptx_path, outdir):
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(outdir),
                    str(pptx_path)], check=True, stdout=subprocess.DEVNULL,
                   stderr=subprocess.DEVNULL, timeout=600)
    return Path(outdir) / (Path(pptx_path).stem + ".pdf")


def words(txt):
    return len(re.sub(r"\[click\]", " ", txt).split())


def section_name(key):
    return D.SLIDES[key].get("label") or {"cover": "Cover", "thanks": "Thank you"}[key]


def fmt_time(sec):
    return f"{sec // 60}:{sec % 60:02d}"


def write_guide():
    rows = []
    tot_t = tot_w = 0
    cum = 0
    for i, key in enumerate(MAIN, 1):
        d = D.SLIDES[key]
        w = words(d["notes"])
        est = round(w / WPM * 60)
        cum += est
        rows.append((i, section_name(key), d["seconds"], w, est, cum))
        tot_t += d["seconds"]
        tot_w += w
    est_total = round(tot_w / WPM * 60)
    L = ["# Cheapify — Presenter Guide", "",
         "Final pitch for the GIZ Youth Entrepreneurship Program. Hard limit: under 6 minutes.", "",
         "## How to use this guide", "",
         "- **Say** is the exact script for each slide; `[click]` marks an animation click, and the "
         "first animation of each slide plays automatically when the slide appears.",
         "- **Key numbers** and **How it's calculated** are your cheat sheet if you are asked where a "
         "figure comes from.",
         "- **If they ask** has short, ready answers for likely questions; the appendix slides are "
         "backups to show only when asked.", "",
         "## Timing", "",
         "| Slide | Section | Target (s) | Words | Est. at 140 wpm (s) | Cumulative (est.) |",
         "| :--- | :--- | :--- | :--- | :--- | :--- |"]
    for r_ in rows:
        L.append(f"| {r_[0]} | {r_[1]} | {r_[2]} | {r_[3]} | {r_[4]} | {fmt_time(r_[5])} |")
    L.append(f"| | **Total** | **{tot_t}** | **{tot_w}** | **{est_total}** | **{fmt_time(est_total)}** |")
    L += ["", f"Estimated delivery time at 140 words per minute: {fmt_time(est_total)} "
              f"(target budget {fmt_time(tot_t)}).", ""]
    for i, key in enumerate(MAIN, 1):
        d = D.SLIDES[key]
        L += [f"## Slide {i} — {section_name(key)}", "", "### Say", "", f"*Time: ~{d['seconds']}s*", "",
              d["notes"], "", "### Key numbers", ""]
        if d["keynums"]:
            L += ["| Number | What it means | Source |", "| :--- | :--- | :--- |"]
            for n, m, src in d["keynums"]:
                L.append(f"| {n} | {m} | {D.SRC[src]} [{src}] |")
        else:
            L.append("No figures on this slide.")
        L += ["", "### How it's calculated", ""]
        if d["calc"]:
            L += [f"- {c}" for c in d["calc"]]
        else:
            L.append("No derived numbers on this slide.")
        L += ["", "### If they ask", ""]
        if d["qa"]:
            for q, a in d["qa"]:
                L += [f"- **Q: {q}** A: {a}"]
        else:
            L.append("No expected questions for this slide.")
        L.append("")
    L += ["## Appendix slides", ""]
    for k in APPX:
        L.append(f"- **{k} — {D.APPENDIX[k]['headline']}:** {D.APPENDIX_WHEN[k]}")
    L += ["", "## References", ""]
    L += [f"{i}. {ref}" for i, ref in enumerate(D.REFERENCES, 1)]
    OUT_GUIDE.write_text("\n".join(L).rstrip() + "\n")
    return rows, tot_t, tot_w


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--static", action="store_true", help="build only the static PPTX")
    ap.add_argument("--qa-states", action="store_true")
    ap.add_argument("--qa-dir", default=str(ROOT / "tmp" / "deck-qa"))
    a = ap.parse_args()
    qa = Path(a.qa_dir)
    qa.mkdir(parents=True, exist_ok=True)
    if a.static:
        out = qa / ("Cheapify-GIZ-Pitch-qastates.pptx" if a.qa_states else "Cheapify-GIZ-Pitch-static.pptx")
        build(static=True, qa_states=a.qa_states, out=out)
        print(out)
        return
    build(static=False, out=OUT_PPTX)
    st = qa / "Cheapify-GIZ-Pitch-static.pptx"
    build(static=True, qa_states=a.qa_states, out=st)
    pdf = convert_pdf(st, qa)
    shutil.copy(pdf, OUT_PDF)
    rows, tt, tw = write_guide()
    print(OUT_PPTX, OUT_PDF, OUT_GUIDE, f"target={tt}s words={tw}")


if __name__ == "__main__":
    main()
